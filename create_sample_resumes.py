
import pymupdf
from docx import Document


resume_text = """
Alex Morgan
Python Developer

Skills: Python, SQL, FastAPI
Experience: Built data processing applications.
"""


# Create a sample PDF.
pdf = pymupdf.open()
page = pdf.new_page()
page.insert_text((72, 72), resume_text)
pdf.save("sample_resume.pdf")
pdf.close()


# Create a sample Word document.
document = Document()
document.add_heading("Alex Morgan", level=1)
document.add_paragraph("Python Developer")
document.add_paragraph("Skills: Python, SQL, FastAPI")
document.add_paragraph("Experience: Built data processing applications.")
document.save("sample_resume.docx")

print("Sample resumes created successfully.")
